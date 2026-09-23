"""
document_processor.py
Handles multi-format document loading, parsing, and text chunking.
Supports: PDF, Word (.docx), Plain Text (.txt, .md), CSV, Excel (.xlsx, .xls), JSON.
"""

import io
import json
import os
import re
import tempfile
from typing import List
from urllib.parse import urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

# Supported file extensions
SUPPORTED_EXTENSIONS = [
    "pdf",
    "docx",
    "txt",
    "md",
    "csv",
    "xlsx",
    "xls",
    "json",
]


def _process_pdf(uploaded_file, chunk_size: int, chunk_overlap: int) -> List[Document]:
    """Extract and chunk text from a PDF file."""
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_file:
        tmp_file.write(uploaded_file.getvalue())
        tmp_path = tmp_file.name

    try:
        loader = PyPDFLoader(tmp_path)
        pages = loader.load()

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

        all_chunks = []
        for page_idx, page in enumerate(pages):
            page_num = page.metadata.get("page", page_idx)
            page_chunks = splitter.split_documents([page])
            for c_idx, chunk in enumerate(page_chunks):
                chunk.metadata["source_filename"] = uploaded_file.name
                chunk.metadata["file_type"] = "pdf"
                chunk.metadata["page"] = page_num + 1 if isinstance(page_num, int) else page_num
                chunk.metadata["location"] = f"Page {chunk.metadata['page']}"
                all_chunks.append(chunk)

        return all_chunks
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


def _process_docx(uploaded_file, chunk_size: int, chunk_overlap: int) -> List[Document]:
    """Extract paragraphs and tables from a Word (.docx) document."""
    import docx

    doc = docx.Document(io.BytesIO(uploaded_file.getvalue()))
    extracted_text_blocks = []

    # 1. Extract paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            extracted_text_blocks.append(text)

    # 2. Extract tables in Markdown table format to preserve columns & rows
    for table_idx, table in enumerate(doc.tables, 1):
        table_lines = [f"\n[Table {table_idx}]"]
        for row_idx, row in enumerate(table.rows):
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            table_lines.append("| " + " | ".join(row_cells) + " |")
            if row_idx == 0:
                table_lines.append("| " + " | ".join(["---"] * len(row_cells)) + " |")
        extracted_text_blocks.append("\n".join(table_lines))

    full_text = "\n\n".join(extracted_text_blocks)
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    raw_chunks = splitter.split_text(full_text)

    documents = []
    for i, chunk_text in enumerate(raw_chunks, 1):
        doc = Document(
            page_content=chunk_text,
            metadata={
                "source_filename": uploaded_file.name,
                "file_type": "docx",
                "page": i,
                "location": f"Section {i}",
            },
        )
        documents.append(doc)
    return documents


def _process_text(uploaded_file, chunk_size: int, chunk_overlap: int) -> List[Document]:
    """Extract and chunk plain text or Markdown files."""
    raw_bytes = uploaded_file.getvalue()
    try:
        content = raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        content = raw_bytes.decode("latin-1", errors="replace")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    raw_chunks = splitter.split_text(content)

    documents = []
    for i, chunk_text in enumerate(raw_chunks, 1):
        doc = Document(
            page_content=chunk_text,
            metadata={
                "source_filename": uploaded_file.name,
                "file_type": "txt",
                "page": i,
                "location": f"Section {i}",
            },
        )
        documents.append(doc)
    return documents


def _process_csv(uploaded_file, rows_per_chunk: int = 15) -> List[Document]:
    """
    Extract and structure CSV tabular data into record-based chunks.
    Preserves column headers with each row so embeddings capture tabular context.
    """
    try:
        df = pd.read_csv(io.BytesIO(uploaded_file.getvalue()))
    except Exception:
        df = pd.read_csv(io.BytesIO(uploaded_file.getvalue()), encoding="latin-1")

    df = df.fillna("N/A")
    total_rows = len(df)
    columns = list(df.columns)

    documents = []
    # Add a schema overview chunk
    overview_text = (
        f"CSV File Overview: {uploaded_file.name}\n"
        f"Total Records: {total_rows}\n"
        f"Columns: {', '.join(str(c) for c in columns)}\n"
        f"Sample Columns & Types: {', '.join(f'{col} ({df[col].dtype})' for col in columns[:8])}"
    )
    documents.append(
        Document(
            page_content=overview_text,
            metadata={
                "source_filename": uploaded_file.name,
                "file_type": "csv",
                "page": 1,
                "location": "Dataset Summary",
            },
        )
    )

    # Group rows into manageable chunks
    for start_idx in range(0, total_rows, rows_per_chunk):
        end_idx = min(start_idx + rows_per_chunk, total_rows)
        chunk_rows = df.iloc[start_idx:end_idx]

        records_text = []
        for row_num, (_, row) in enumerate(chunk_rows.iterrows(), start=start_idx + 1):
            row_items = [f"{col}: {row[col]}" for col in columns]
            records_text.append(f"[Record {row_num}] " + " | ".join(row_items))

        chunk_content = (
            f"Table: {uploaded_file.name} (Rows {start_idx + 1} to {end_idx} of {total_rows})\n"
            + "\n".join(records_text)
        )

        doc = Document(
            page_content=chunk_content,
            metadata={
                "source_filename": uploaded_file.name,
                "file_type": "csv",
                "page": (start_idx // rows_per_chunk) + 2,
                "location": f"Rows {start_idx + 1}-{end_idx}",
            },
        )
        documents.append(doc)

    return documents


def _process_excel(uploaded_file, rows_per_chunk: int = 15) -> List[Document]:
    """Extract all sheets from Excel spreadsheets (.xlsx, .xls) and format as tabular records."""
    excel_file = pd.ExcelFile(io.BytesIO(uploaded_file.getvalue()))
    documents = []
    page_counter = 1

    for sheet_name in excel_file.sheet_names:
        df = pd.read_excel(excel_file, sheet_name=sheet_name).fillna("N/A")
        total_rows = len(df)
        columns = list(df.columns)

        # Sheet overview chunk
        overview_text = (
            f"Excel Sheet: '{sheet_name}' in {uploaded_file.name}\n"
            f"Total Rows: {total_rows} | Columns: {', '.join(str(c) for c in columns)}"
        )
        documents.append(
            Document(
                page_content=overview_text,
                metadata={
                    "source_filename": uploaded_file.name,
                    "file_type": "excel",
                    "page": page_counter,
                    "location": f"Sheet '{sheet_name}' Summary",
                },
            )
        )
        page_counter += 1

        if total_rows == 0:
            continue

        for start_idx in range(0, total_rows, rows_per_chunk):
            end_idx = min(start_idx + rows_per_chunk, total_rows)
            chunk_rows = df.iloc[start_idx:end_idx]

            records_text = []
            for row_num, (_, row) in enumerate(chunk_rows.iterrows(), start=start_idx + 1):
                row_items = [f"{col}: {row[col]}" for col in columns]
                records_text.append(f"[Row {row_num}] " + " | ".join(row_items))

            chunk_content = (
                f"Sheet '{sheet_name}' (Rows {start_idx + 1} to {end_idx} of {total_rows})\n"
                + "\n".join(records_text)
            )

            doc = Document(
                page_content=chunk_content,
                metadata={
                    "source_filename": uploaded_file.name,
                    "file_type": "excel",
                    "page": page_counter,
                    "location": f"Sheet '{sheet_name}', Rows {start_idx + 1}-{end_idx}",
                },
            )
            documents.append(doc)
            page_counter += 1

    return documents


def _process_json(uploaded_file, chunk_size: int = 1000, chunk_overlap: int = 200) -> List[Document]:
    """Parse JSON data (arrays of objects or nested dictionaries)."""
    raw_bytes = uploaded_file.getvalue()
    try:
        data = json.loads(raw_bytes.decode("utf-8"))
    except UnicodeDecodeError:
        data = json.loads(raw_bytes.decode("latin-1", errors="replace"))

    documents = []

    # If it's a list of objects (e.g. products, users, restaurant items)
    if isinstance(data, list):
        total_records = len(data)
        batch_size = 5
        for start_idx in range(0, total_records, batch_size):
            end_idx = min(start_idx + batch_size, total_records)
            batch = data[start_idx:end_idx]

            items_formatted = []
            for idx, item in enumerate(batch, start=start_idx + 1):
                if isinstance(item, dict):
                    formatted_props = [f"{k}: {v}" for k, v in item.items()]
                    items_formatted.append(f"[Item {idx}]\n" + "\n".join(formatted_props))
                else:
                    items_formatted.append(f"[Item {idx}]: {item}")

            chunk_content = (
                f"JSON Data: {uploaded_file.name} (Items {start_idx + 1} to {end_idx} of {total_records})\n\n"
                + "\n\n---\n\n".join(items_formatted)
            )

            page_num = (start_idx // batch_size) + 1
            doc = Document(
                page_content=chunk_content,
                metadata={
                    "source_filename": uploaded_file.name,
                    "file_type": "json",
                    "page": page_num,
                    "location": f"Items {start_idx + 1}-{end_idx}",
                },
            )
            documents.append(doc)

    # If it's a nested dictionary / single object
    elif isinstance(data, dict):
        formatted_text = json.dumps(data, indent=2, ensure_ascii=False)
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", ", ", " ", ""],
        )
        raw_chunks = splitter.split_text(formatted_text)
        for i, chunk_text in enumerate(raw_chunks, 1):
            doc = Document(
                page_content=f"JSON Document: {uploaded_file.name} (Part {i})\n\n{chunk_text}",
                metadata={
                    "source_filename": uploaded_file.name,
                    "file_type": "json",
                    "page": i,
                    "location": f"Section {i}",
                },
            )
            documents.append(doc)
    else:
        # Fallback scalar or simple value
        doc = Document(
            page_content=str(data),
            metadata={
                "source_filename": uploaded_file.name,
                "file_type": "json",
                "page": 1,
                "location": "Data Root",
            },
        )
        documents.append(doc)

    return documents


def load_and_split_document(
    uploaded_file, chunk_size: int = 1000, chunk_overlap: int = 200
) -> List[Document]:
    """
    Main entry point: Automatically detects file extension and routes
    to the optimal parser for PDF, DOCX, TXT, CSV, Excel, or JSON.

    Args:
        uploaded_file: Streamlit UploadedFile object.
        chunk_size: Target characters per chunk.
        chunk_overlap: Overlap characters between chunks.

    Returns:
        List of LangChain Document objects with rich metadata.
    """
    filename = uploaded_file.name.lower()
    ext = os.path.splitext(filename)[1].lstrip(".")

    if ext == "pdf":
        return _process_pdf(uploaded_file, chunk_size, chunk_overlap)
    elif ext == "docx":
        return _process_docx(uploaded_file, chunk_size, chunk_overlap)
    elif ext in ["txt", "md"]:
        return _process_text(uploaded_file, chunk_size, chunk_overlap)
    elif ext == "csv":
        return _process_csv(uploaded_file)
    elif ext in ["xlsx", "xls"]:
        return _process_excel(uploaded_file)
    elif ext == "json":
        return _process_json(uploaded_file, chunk_size, chunk_overlap)
    else:
        # Fallback to plain text processor
        return _process_text(uploaded_file, chunk_size, chunk_overlap)


def load_and_split_url(url: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> List[Document]:
    """
    Fetch a web page, extract main content text, and split into chunks.

    Args:
        url: Full HTTP or HTTPS web URL.
        chunk_size: Target characters per chunk.
        chunk_overlap: Overlap characters between chunks.

    Returns:
        List of LangChain Document objects with rich metadata.
    """
    url = url.strip()
    if not url.startswith("http://") and not url.startswith("https://"):
        url = "https://" + url

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/120.0.0.0 Safari/537.36"
        )
    }

    resp = requests.get(url, headers=headers, timeout=15)
    resp.raise_for_status()

    soup = BeautifulSoup(resp.text, "html.parser")

    # Extract title
    title = ""
    if soup.title and soup.title.string:
        title = soup.title.string.strip()
    if not title:
        parsed = urlparse(url)
        title = parsed.netloc + (parsed.path if parsed.path and parsed.path != "/" else "")

    # Clean out non-content elements
    for tag in soup(["script", "style", "nav", "header", "footer", "aside", "noscript", "svg", "iframe", "form"]):
        tag.decompose()

    # Extract text from main container or body
    main_elem = soup.find("main") or soup.find("article") or soup.find("div", {"id": "content"}) or soup.find("div", {"class": "content"}) or soup.body
    if main_elem:
        raw_text = main_elem.get_text(separator="\n", strip=True)
    else:
        raw_text = soup.get_text(separator="\n", strip=True)

    # Clean excessive blank lines
    cleaned_lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    content_text = "\n\n".join(cleaned_lines)

    if not content_text:
        raise ValueError(f"Could not extract readable text from {url}")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    raw_chunks = splitter.split_text(content_text)
    documents = []
    display_title = title[:60] if len(title) > 60 else title

    for idx, chunk_text in enumerate(raw_chunks, 1):
        doc = Document(
            page_content=f"Web Source: {display_title}\nURL: {url}\n\n{chunk_text}",
            metadata={
                "source_filename": display_title,
                "source_url": url,
                "file_type": "web",
                "page": idx,
                "location": f"Section {idx}",
                "title": display_title,
            },
        )
        documents.append(doc)

    return documents


# Alias for backward compatibility
def load_and_split_pdf(uploaded_file, chunk_size: int = 1000, chunk_overlap: int = 200) -> List[Document]:
    """Backward-compatible alias for load_and_split_document."""
    return load_and_split_document(uploaded_file, chunk_size, chunk_overlap)


def get_document_stats(chunks: List[Document]) -> dict:
    """Return comprehensive metadata and statistics about the processed document."""
    if not chunks:
        return {}

    locations = set(chunk.metadata.get("location", "") for chunk in chunks)
    file_type = chunks[0].metadata.get("file_type", "document") if chunks else "document"
    total_chars = sum(len(chunk.page_content) for chunk in chunks)
    source_url = chunks[0].metadata.get("source_url", "") if chunks else ""

    unit_mapping = {
        "pdf": "Pages",
        "docx": "Sections",
        "txt": "Sections",
        "csv": "Row Blocks",
        "excel": "Sheets/Blocks",
        "json": "Item Blocks",
        "web": "Web Sections",
    }
    unit_label = unit_mapping.get(file_type, "Sections")

    return {
        "total_chunks": len(chunks),
        "total_units": len(locations),
        "unit_label": unit_label,
        "file_type": file_type.upper(),
        "total_characters": total_chars,
        "avg_chunk_size": total_chars // len(chunks) if chunks else 0,
        "source_url": source_url,
    }

